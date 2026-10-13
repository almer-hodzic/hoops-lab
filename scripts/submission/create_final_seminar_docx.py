#!/usr/bin/env python3
"""Create the final Bosnian seminar paper DOCX from verified submission assets."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SUB = ROOT / "submission"
DOCS = SUB / "docs"
ASSETS = SUB / "report_assets"
OUT = DOCS / "NBA_Clutch_Agent_Seminarski.docx"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


METRICS = load_json(SUB / "metrics_summary.json")
DATASET = METRICS["dataset"]
PROD = METRICS["production_model"]["test_regular_2025_26"]
VAL = METRICS["production_model"]["validation_regular"]
CLUTCH = METRICS["clutch_performance"]
DECISION = METRICS["decision_policy"]
DEFENSE = METRICS["defensive_experiment"]["selected_v11"]
DEFENSE_CSV = pd.read_csv(SUB / "report_data" / "defensive_shot_quality_model_metrics.csv")
MODEL_CMP = pd.read_csv(SUB / "report_data" / "model_comparison_test_regular.csv")


def fmt_int(n: int | float) -> str:
    return f"{int(n):,}".replace(",", ".")


def fmt(x: float, digits: int = 4) -> str:
    return f"{float(x):.{digits}f}"


def pct(x: float, digits: int = 2) -> str:
    return f"{float(x) * 100:.{digits}f}%"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(10)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table, header_fill: str = "1F4E79") -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_border(cell)
            cell.margin_top = 80
            cell.margin_bottom = 80
            cell.margin_left = 80
            cell.margin_right = 80
            if r == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.bold = True
            elif r % 2 == 0:
                set_cell_shading(cell, "F5F8FC")


def add_table(doc: Document, title: str, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    cap = doc.add_paragraph()
    cap.style = "Caption"
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(title).bold = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
        if widths:
            hdr[i].width = Cm(widths[i])
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], str(value), align=align)
            if widths:
                cells[i].width = Cm(widths[i])
    style_table(table)
    doc.add_paragraph()


def add_figure(doc: Document, filename: str, caption: str, analysis: list[str], width: float = 5.9) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(ASSETS / filename), width=Inches(width))
    cap = doc.add_paragraph()
    cap.style = "Caption"
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    cap.add_run(caption).bold = True
    for text in analysis:
        add_body(doc, text)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = "Body Text"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.7)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    p.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_h1(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=1)
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True


def add_h2(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=2)
    p.paragraph_format.keep_with_next = True


def add_page_number(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    p._p.append(fld)


def restart_page_numbering(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def apply_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)

    body = doc.styles["Body Text"]
    body.font.name = "Times New Roman"
    body.font.size = Pt(12)

    for style_name, size in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)

    cap = doc.styles["Caption"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(10)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor(0, 0, 0)


def title_page(doc: Document) -> None:
    for _ in range(1):
        doc.add_paragraph()
    top = doc.add_paragraph()
    top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in ['UNIVERZITET "DŽEMAL BIJEDIĆ" U MOSTARU', "FAKULTET INFORMACIJSKIH TEHNOLOGIJA"]:
        r = top.add_run(line + "\n")
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.bold = True
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SEMINARSKI RAD")
    r.font.size = Pt(16)
    r.bold = True
    for _ in range(1):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("NBA CLUTCH AGENT")
    r.font.size = Pt(22)
    r.bold = True
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Predikcija kvaliteta NBA šuta i podrška odlučivanju u clutch situacijama")
    r.font.size = Pt(14)
    for _ in range(3):
        doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    info.paragraph_format.left_indent = Cm(7.0)
    for line in [
        "Student: Almer Hodžić",
        "Broj indeksa: [BROJ INDEKSA]",
        "Predmet: [NAZIV PREDMETA]",
        "Profesor/mentor: [IME PROFESORA]",
    ]:
        rr = info.add_run(line + "\n")
        rr.font.size = Pt(12)
    for _ in range(3):
        doc.add_paragraph()
    bottom = doc.add_paragraph()
    bottom.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bottom.add_run("Mostar, 2026.").font.size = Pt(12)


def add_static_toc(doc: Document) -> None:
    doc.add_heading("Sadržaj", level=1)
    entries = [
        "1. O projektu i motivacija",
        "2. Izvor i nastanak dataseta",
        "3. Struktura podataka i ciljna varijabla",
        "4. Eksplorativna analiza podataka",
        "5. Priprema podataka i sprječavanje curenja informacija",
        "6. Izbor metoda i evaluacijskih metrika",
        "7. Treniranje modela i razvoj kroz iteracije",
        "8. Evaluacija finalnog Shot Quality modela",
        "9. Usporedba modela",
        "10. Decision Policy: TAKE / REVIEW / DEFER",
        "11. Eksperimentalni kontekst odbrane",
        "12. Integracija modela u aplikaciju",
        "13. Ograničenja i mogućnosti poboljšanja",
        "14. Zaključak",
        "15. Literatura i izvori",
    ]
    for e in entries:
        p = doc.add_paragraph()
        p.style = "Body Text"
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(e)
    doc.add_paragraph()


def build_doc() -> None:
    doc = Document()
    apply_styles(doc)
    title_page(doc)

    doc.add_section(WD_SECTION.NEW_PAGE)
    for section in doc.sections[1:]:
        section.footer.is_linked_to_previous = False
        restart_page_numbering(section, 1)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        add_page_number(section)

    add_static_toc(doc)

    add_h1(doc, "1. O projektu i motivacija")
    add_body(doc, "NBA Clutch Agent je akademski projekat iz oblasti mašinskog učenja koji povezuje analizu sportskih podataka, vjerovatnosno modeliranje i donošenje odluka. Osnovni zadatak nije samo klasificirati šut kao pogodak ili promašaj, nego procijeniti vjerovatnoću pogotka prije nego što je šut izveden. Ta vjerovatnoća se zatim prevodi u očekivane poene po šutu i koristi kao ulaz u jednostavan, transparentan sloj odlučivanja.")
    add_body(doc, "Motivacija dolazi iz činjenice da u košarci kvalitet odluke nije uvijek isti kao rezultat pojedinačnog šuta. Dobar šut može biti promašen, a loš šut može slučajno biti pogođen. Zbog toga je za akademsku analizu korisnije procijeniti kvalitet pokušaja na osnovu informacija koje su bile dostupne prije šuta: lokacija, udaljenost, tip šuta, vrijeme, rezultat, historijski profil šutera i kontekst utakmice.")
    add_body(doc, "Projekat je jači od jednostavne vježbe binarne klasifikacije jer uključuje kompletan tok: prikupljanje i organizaciju NBA event podataka, rekonstrukciju pre-shot konteksta, temporalnu validaciju, sprječavanje curenja informacija, kalibraciju vjerovatnoća, decision policy i web aplikaciju za demonstraciju. Posebno je važno što produkcijski model i eksperimentalna istraživanja nisu pomiješani: produkcijski Shot Quality model daje odluke, dok je defanzivni tracking kontekst jasno označen kao istraživačka ekstenzija.")
    add_body(doc, "Konačni cilj rada je pokazati kako se mašinsko učenje može koristiti kao podrška odlučivanju u realističnom sportskom problemu. Model ne tvrdi da može garantovati ishod šuta i ne tvrdi kauzalni efekat odluke. On procjenjuje vjerovatnoću i omogućava poređenje očekivane vrijednosti šuta sa transparentnim replacement benchmarkom.")

    add_h1(doc, "2. Izvor i nastanak dataseta")
    add_body(doc, f"Primarni dataset je `data/processed/shots_enriched.parquet`. On sadrži {fmt_int(DATASET['rows'])} pokušaja šuta iz {fmt_int(DATASET['games'])} utakmica, za sezone {', '.join(DATASET['seasons'])}. Podaci obuhvataju regularnu sezonu i playoff, sa datumskim rasponom od {DATASET['date_min']} do {DATASET['date_max']}.")
    add_body(doc, "Dataset je nastao kombinovanjem NBA Shot Detail arhiva i PlayByPlay V3 podataka. Shot Detail daje osnovne informacije o šutu: lokaciju, udaljenost, zonu, tip akcije i rezultat. PlayByPlay podaci su korišteni za provjeru i rekonstrukciju konteksta prije šuta, naročito rezultata utakmice i clutch situacija.")
    add_body(doc, "Kao primarni akademski problem izabran je moderni canonical dataset iz sezona 2020-21 do 2025-26. Stari 2014/15 tracking dataset nije primarni dataset rada; on je obrađen kao dodatni eksperiment za defanzivni kontekst. Takva podjela je važna jer moderni produkcijski model i eksperimentalni tracking model nemaju isti izvor, vremenski period ni isti nivo pouzdanosti za savremenu primjenu.")
    add_table(
        doc,
        "Tabela A. Pregled primarnog dataseta",
        ["Element", "Vrijednost"],
        [
            ["Ukupan broj pokušaja", fmt_int(DATASET["rows"])],
            ["Broj utakmica", fmt_int(DATASET["games"])],
            ["Sezone", ", ".join(DATASET["seasons"])],
            ["Regularna sezona", fmt_int(DATASET["regular_season_rows"])],
            ["Playoff", fmt_int(DATASET["playoff_rows"])],
            ["Clutch pokušaji", fmt_int(DATASET["clutch_rows"])],
            ["Jedinstveni šuteri", fmt_int(DATASET["unique_shooters"])],
            ["Pogođeno / promašeno", f"{fmt_int(DATASET['target']['made'])} / {fmt_int(DATASET['target']['missed'])}"],
            ["2PT / 3PT pokušaji", f"{fmt_int(DATASET['shot_value_distribution']['2'])} / {fmt_int(DATASET['shot_value_distribution']['3'])}"],
        ],
        widths=[6, 9],
    )
    add_body(doc, "Clutch definicija u projektu je službena radna definicija: četvrta četvrtina ili produžetak, pet minuta ili manje do kraja perioda, i apsolutna pre-shot razlika u rezultatu najviše pet poena. Ova definicija omogućava dosljedno poređenje regularnih i clutch situacija bez ručnog biranja primjera.")

    add_h1(doc, "3. Struktura podataka i ciljna varijabla")
    add_body(doc, "Ciljna varijabla je `shot_made`, odnosno binarna oznaka da li je šut pogođen. U cijelom canonical datasetu pogođeno je " + fmt_int(DATASET["target"]["made"]) + " šuteva, a promašeno " + fmt_int(DATASET["target"]["missed"]) + f". Ukupni FG% iznosi {pct(DATASET['target']['made_rate'])}.")
    add_body(doc, "Pre-shot ulazi mogu se grupisati u nekoliko skupina: identifikacijski i kontekstualni podaci, geometrija šuta, stanje utakmice, zone i kategorije šuta, te leakage-safe historijski profil šutera. U produkcijskom modelu ne koriste se post-outcome polja kao ulaz, jer bi to direktno unijelo curenje informacija.")
    add_body(doc, f"Dataset sadrži {fmt_int(DATASET['shot_value_distribution']['2'])} pokušaja za dva poena i {fmt_int(DATASET['shot_value_distribution']['3'])} pokušaja za tri poena. Uspješnost po vrijednosti šuta je {pct(DATASET['shot_value_success']['2'])} za 2PT i {pct(DATASET['shot_value_success']['3'])} za 3PT. To potvrđuje poznatu košarkašku razliku: šutevi za tri poena su teži, ali svaki pogodak vrijedi više.")
    add_figure(
        doc,
        "01_target_distribution.png",
        "Slika 1. Distribucija ciljne varijable: pogođeni i promašeni šutevi.",
        [
            "Slika prikazuje da dataset nije ekstremno neuravnotežen, ali promašaji ipak čine većinu pokušaja. Ova činjenica je važna jer model ne smije jednostavno naučiti većinsku klasu; cilj je procijeniti vjerovatnoću za svaki pojedinačni pokušaj.",
            "Za ovaj problem sama tačnost nije dovoljna. Ako model predvidi 47% za veliki broj šuteva, on može biti koristan iako se pojedinačni ishodi često razlikuju. Zato su Log Loss, Brier Score i kalibracija centralne metrike."
        ],
    )
    add_figure(
        doc,
        "03_shot_value_and_zone_success.png",
        "Slika 2. Uspješnost šuta po vrijednosti i glavnim zonama.",
        [
            "Uspješnost jasno zavisi od geometrije i zone. Bliski šutevi obično imaju veću vjerovatnoću pogotka, dok daleki šutevi imaju niži FG%, ali mogu imati veću očekivanu vrijednost zbog tri poena.",
            "Ovaj obrazac opravdava korištenje `shot_distance`, koordinata, zone i tipa šuta kao važnih ulaza u model."
        ],
    )

    add_h1(doc, "4. Eksplorativna analiza podataka")
    add_body(doc, "Eksplorativna analiza ima dvije uloge. Prvo, ona pomaže razumjeti osnovne obrasce u košarkaškim podacima. Drugo, ona provjerava da li su ulazi i target varijabla konzistentni sa očekivanim ponašanjem igre. Ako bi, na primjer, daleki šutevi imali viši FG% od bliskih bez objašnjenja, to bi bio signal za dodatni audit podataka.")
    add_figure(
        doc,
        "02_shot_distance_distribution.png",
        "Slika 3. Distribucija udaljenosti šuta.",
        [
            "Distribucija udaljenosti pokazuje koncentraciju pokušaja blizu obruča i oko trice. To je u skladu sa savremenom NBA igrom, u kojoj su najčešće zone obruč i linija za tri poena.",
            "Za model je ova distribucija bitna jer neke udaljenosti imaju mnogo veću podršku u podacima od drugih. Predikcije u dobro zastupljenim zonama obično su stabilnije od predikcija u rijetkim kombinacijama akcije, distance i konteksta."
        ],
    )
    add_figure(
        doc,
        "04_clutch_vs_non_clutch.png",
        "Slika 4. Poređenje clutch i non-clutch konteksta.",
        [
            "Clutch subset je mnogo manji od ostatka dataseta i ima specifičan kontekst igre: vrijeme je ograničeno, rezultat je blizu, a selekcija šuteva može biti drugačija. Slika ne dokazuje kauzalnost, nego opisuje razliku u historijskim grupama.",
            "Zbog manjeg uzorka, clutch evaluacija se u radu posmatra kao posebno važan, ali statistički osjetljiv segment. Model mora zadržati diskriminativnost i kalibraciju i u takvim situacijama."
        ],
    )
    add_body(doc, "EDA također potvrđuje potrebu za posebnim tretmanom playoffa. Playoff pokušaji imaju drugačiji kontekst selekcije, protivnika i intenziteta igre, zbog čega se playoff performanse ne smiju nekritički miješati sa regularnom sezonom pri model selection procesu.")

    add_h1(doc, "5. Priprema podataka i sprječavanje curenja informacija")
    add_body(doc, "Najvažniji metodološki princip u projektu je da model smije koristiti samo informacije dostupne prije šuta. To znači da target, poeni nakon šuta, kasniji šutevi istog igrača i budući rezultati utakmica ne smiju ulaziti u feature engineering za trenutni red.")
    add_body(doc, "Podaci su uređeni hronološki. Trening koristi regularne sezone 2020-21 do 2023-24, validacija koristi sezonu 2024-25, a finalni test je sezona 2025-26. Ovakav temporalni split je bliži stvarnoj upotrebi modela: u praksi se model trenira na prošlosti i primjenjuje na budućnost.")
    add_table(
        doc,
        "Tabela B. Temporalni split za produkcijski model",
        ["Split", "Sezone", "Namjena"],
        [
            ["Train", "2020-21, 2021-22, 2022-23, 2023-24", "Učenje modela i parametara"],
            ["Validation", "2024-25", "Model selection i kalibracija bez korištenja test sezone"],
            ["Test", "2025-26", "Finalna held-out evaluacija"],
        ],
        widths=[3, 7, 6],
    )
    add_body(doc, "Nasumični split nije poželjan jer bi u isti trening i test mogli dospjeti redovi iz iste sezone, sličnog perioda forme ili istog konteksta igrača. To bi precijenilo realnu sposobnost generalizacije. Temporalna validacija je strožija i realističnija.")
    add_body(doc, "Historijski profili šutera računaju se samo iz prethodnih pokušaja. Primjeri takvih featurea su prior FG%, prior 2P/3P profil, zona, distanca, season-to-date statistika i previous-season profil. Shrunk/Bayesian verzije procenata ublažavaju problem malih uzoraka, jer igrač sa dva pogođena šuta iz dva pokušaja ne smije automatski izgledati kao savršen šuter.")
    add_body(doc, "Cold-start slučajevi nastaju kada igrač nema dovoljno prethodne historije. U tim slučajevima se koristi populacioni fallback ili ograničeni dostupni profil, a model dobija signal da je profil hladan start. Time se izbjegava izmišljanje precizne historije za igrača koji je nema.")

    add_h1(doc, "6. Izbor metoda i evaluacijskih metrika")
    add_body(doc, "U radu su testirani baseline modeli, Logistic Regression i HistGradientBoosting. Baseline modeli su važni jer daju referentnu tačku: složeniji model ima smisla samo ako jasno poboljšava vjerovatnosne metrike u odnosu na jednostavna pravila.")
    add_body(doc, "Logistic Regression je interpretabilniji linearni model i koristan je kao stabilna početna tačka. HistGradientBoosting može bolje uhvatiti nelinearne odnose, na primjer interakciju distance, akcije, zone, vrijednosti šuta i profila šutera. Konačni izbor nije zasnovan na kompleksnosti, nego na validacijskim i test metrikama.")
    add_body(doc, "Log Loss kažnjava uvjerene pogrešne predikcije i zato je prikladan za vjerovatnosne modele. Brier Score mjeri kvadratnu grešku vjerovatnoće. ROC AUC mjeri rangiranje pozitivnih i negativnih ishoda, dok kalibracija provjerava da li npr. predikcije od 60% zaista približno odgovaraju 60% pogodaka u sličnim grupama.")
    add_body(doc, "Accuracy nije primarna metrika jer bi prag od 0.5 pretvorio probabilistički problem u grubu klasifikaciju. U košarci je razlika između 38%, 43% i 62% vrlo važna, čak i ako bi sve tri vrijednosti nakon hard thresholda mogle završiti u istoj klasi.")

    add_h1(doc, "7. Treniranje modela i razvoj kroz iteracije")
    add_body(doc, "Razvoj modela tekao je postupno. Prvo su korišteni globalni i shot-value baseline modeli. Oni pokazuju koliko se može dobiti bez detaljnog konteksta. Zatim su uvedeni geometrija šuta, tip akcije i zone. Nakon toga su testirani raw player ID i historijski profil šutera.")
    add_body(doc, "Raw player ID može pomoći modelu da nauči razliku između igrača, ali sam po sebi nije idealan za održavanje i cold-start situacije. Historijski profil šutera je bolji akademski feature set jer opisuje šta je igrač uradio prije trenutnog šuta. On je interpretabilniji i direktno povezan sa košarkaškim znanjem.")
    add_body(doc, "U konačnim eksperimentima testirane su konfiguracije sa generičkim featureima, raw player ID, historical shooter profile i kombinacija profila sa raw ID. Sigmoid kalibracija je izabrana jer je dala stabilnu vjerovatnosnu interpretaciju bez korištenja test sezone za izbor modela.")
    add_figure(
        doc,
        "06_model_comparison.png",
        "Slika 5. Usporedba modela prema test Log Loss metriki.",
        [
            "Slika prikazuje da napredniji modeli sa geometrijom, kontekstom i profilom šutera nadmašuju jednostavne baseline pristupe. Razlika u Log Loss metriki je akademski važna jer poboljšava kvalitet vjerovatnoće, a ne samo hard klasifikaciju.",
            "Model nije odabran zato što je najkompleksniji, nego zato što daje najkorisniji balans između performansi, kalibracije i interpretabilnosti."
        ],
    )

    add_h1(doc, "8. Evaluacija finalnog Shot Quality modela")
    add_body(doc, f"Finalni held-out test je 2025-26 regularna sezona. Na tom skupu model ima Log Loss {fmt(PROD['log_loss'], 5)}, Brier Score {fmt(PROD['brier'], 5)}, ROC AUC {fmt(PROD['roc_auc'], 5)}, stvarni FG% {fmt(PROD['actual_fg_pct'], 4)} i predviđeni FG% {fmt(PROD['predicted_fg_pct'], 4)}.")
    add_table(
        doc,
        "Tabela C. Produkcijske metrike finalnog modela",
        ["Skup", "Redova", "Log Loss", "Brier", "ROC AUC", "Stvarni FG%", "Predviđeni FG%"],
        [
            ["Validation 2024-25", fmt_int(VAL["rows"]), fmt(VAL["log_loss"], 5), fmt(VAL["brier"], 5), fmt(VAL["roc_auc"], 5), fmt(VAL["actual_fg_pct"], 4), fmt(VAL["predicted_fg_pct"], 4)],
            ["Test 2025-26", fmt_int(PROD["rows"]), fmt(PROD["log_loss"], 5), fmt(PROD["brier"], 5), fmt(PROD["roc_auc"], 5), fmt(PROD["actual_fg_pct"], 4), fmt(PROD["predicted_fg_pct"], 4)],
            ["Clutch test 2025-26", fmt_int(CLUTCH["rows"]), fmt(CLUTCH["log_loss"], 5), fmt(CLUTCH["brier"], 5), fmt(CLUTCH["roc_auc"], 5), fmt(CLUTCH["actual_fg_pct"], 4), fmt(CLUTCH["predicted_fg_pct"], 4)],
        ],
    )
    add_figure(
        doc,
        "05_calibration_curve.png",
        "Slika 6. Kalibracijska kriva finalnog produkcijskog modela.",
        [
            "Kalibracija je ključna jer Decision Policy koristi numeričku vjerovatnoću kao ulaz u očekivane poene. Ako model sistematski precjenjuje ili potcjenjuje vjerovatnoće, tada i odluke o TAKE/REVIEW/DEFER mogu biti pomjerene.",
            "Idealna dijagonala predstavlja savršenu kalibraciju. Odstupanja su očekivana u realnim podacima, ali cilj je da predviđene vjerovatnoće budu dovoljno pouzdane za poređenje očekivane vrijednosti."
        ],
    )
    add_figure(
        doc,
        "08_probability_distribution.png",
        "Slika 7. Distribucija predviđenih make-probability vrijednosti.",
        [
            "Distribucija predikcija pokazuje da model ne vraća samo jednu prosječnu vrijednost, nego razlikuje lakše i teže pokušaje. To je važno za decision layer, jer očekivana vrijednost direktno zavisi od `p_make`.",
            "Preklapanje između pogođenih i promašenih šuteva je normalno: pojedinačni ishod je stohastičan i zavisi od faktora koji nisu u potpunosti posmatrani."
        ],
    )
    add_figure(
        doc,
        "11_roc_comparison.png",
        "Slika 8. ROC kriva produkcijskog modela i baselinea.",
        [
            "ROC AUC prikazuje sposobnost modela da rangira pokušaje prema vjerovatnoći pogotka. Produkcijski model nadmašuje globalni baseline, što znači da naučeni kontekst nosi informaciju iznad prosječnog FG%.",
            "Ipak, ROC nije jedina metrika jer model nije samo rang lista; njegova vjerovatnoća se dalje koristi u izračunu očekivane vrijednosti."
        ],
    )
    add_figure(
        doc,
        "12_confusion_matrix.png",
        "Slika 9. Dopunska confusion matrix pri pragu 0.5.",
        [
            "Confusion matrix je uključena jer je poznata u akademskim ML izvještajima, ali u ovom projektu ima dopunski karakter. Prag 0.5 nije optimizovan za košarkašku odluku i ne odražava direktno očekivanu vrijednost šuta.",
            "Model je primarno vjerovatnosni model. Zbog toga se zaključci o kvalitetu modela zasnivaju prvenstveno na Log Loss, Brier Score, ROC AUC i kalibraciji."
        ],
    )

    add_h1(doc, "9. Usporedba modela")
    rows = []
    for row in MODEL_CMP.sort_values("log_loss").head(8).itertuples():
        rows.append([row.model, row.feature_config, row.calibration, fmt(row.log_loss, 5), fmt(row.brier, 5), fmt(row.roc_auc, 5)])
    add_table(
        doc,
        "Tabela D. Glavna usporedba modela na 2025-26 test skupu",
        ["Model", "Feature config", "Kalibracija", "Log Loss", "Brier", "ROC AUC"],
        rows,
    )
    add_body(doc, "Tabela pokazuje da su baseline modeli korisni kao referenca, ali ne dostižu performanse modela koji koristi geometriju, kontekst i profil šutera. Historical shooter profile daje posebno važan doprinos jer unosi informaciju o sposobnosti igrača bez korištenja budućih događaja.")
    add_body(doc, "Razlike u Log Loss metriki mogu izgledati numerički male, ali na stotinama hiljada pokušaja one predstavljaju stabilno poboljšanje vjerovatnosne procjene. U decision sistemu i mala sistematska promjena vjerovatnoće može promijeniti marginu između TAKE, REVIEW i DEFER.")
    add_figure(
        doc,
        "07_feature_importance.png",
        "Slika 10. Permutation feature importance za finalni model.",
        [
            "Najvažniji featurei su povezani sa udaljenošću, tipom akcije i historijskim profilom šutera. Ovaj rezultat je košarkaški smislen: lokacija i način šuta snažno utiču na vjerovatnoću, a profil šutera pomaže razlikovati kvalitet izvođača.",
            "Feature importance se tumači kao prediktivna korisnost, ne kao kauzalni dokaz. Na primjer, `action_type` može biti jak signal, ali ne znači da sama promjena naziva akcije uzrokuje promjenu realizacije."
        ],
    )
    add_body(doc, "Posebno je korisno što se među najvažnijim signalima pojavljuju i osnovni košarkaški faktori i historijske osobine igrača. To pokazuje da model ne zavisi samo od identiteta šutera, nego kombinuje kvalitet lokacije, način izvođenja i prethodno demonstriranu šutersku sposobnost.")
    add_body(doc, "U akademskom smislu ova usporedba pokazuje razliku između modela koji samo pamti prosjek i modela koji koristi strukturirane informacije. Finalni izbor nije tretiran kao apsolutno najbolji mogući NBA model, nego kao najodbranjiviji model u okviru dostupnih podataka, vremenskog splita i potrebe za kalibrisanom vjerovatnoćom.")

    add_h1(doc, "10. Decision Policy: TAKE / REVIEW / DEFER")
    add_body(doc, "Shot Quality model daje `p_make`, ali aplikaciji je potrebna odluka koja je razumljiva korisniku. Decision Policy zato prevodi vjerovatnoću u očekivane poene po šutu i poredi ih sa replacement benchmarkom.")
    add_body(doc, "Osnovne formule su:")
    eq1 = doc.add_paragraph()
    eq1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eq1.add_run("shot_xPPS = p_make × shot_value").bold = True
    eq2 = doc.add_paragraph()
    eq2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eq2.add_run("decision_margin = shot_xPPS - replacement_xPPS").bold = True
    add_body(doc, f"U finalnoj politici replacement xPPS iznosi približno {fmt(DECISION['replacement_xpps'], 6)}, a prag odluke je ±{fmt(DECISION['threshold'], 2)}. Ako je margina veća od +0.15, preporuka je TAKE. Ako je manja od -0.15, preporuka je DEFER. Svi ostali slučajevi su REVIEW.")
    add_body(doc, "Ova pravila približno odgovaraju sljedećim granicama: za 2PT šut DEFER je ispod oko 46.5%, REVIEW oko 46.5%-61.5%, a TAKE iznad oko 61.5%; za 3PT šut DEFER je ispod oko 31.0%, REVIEW oko 31.0%-41.0%, a TAKE iznad oko 41.0%.")
    dec_rows = []
    for row in DECISION["test_regular_2025_26"]["decisions"]:
        dec_rows.append([row["decision"], fmt_int(row["rows"]), pct(row["share"]), fmt(row["mean_shot_xpps"], 3), fmt(row["observed_pps"], 3), fmt(row["mean_margin"], 3)])
    add_table(
        doc,
        "Tabela E. Decision Policy na 2025-26 regular season test skupu",
        ["Odluka", "Redova", "Udio", "Srednji shot xPPS", "Observed PPS", "Srednja margina"],
        dec_rows,
    )
    add_figure(
        doc,
        "09_decision_policy_analysis.png",
        "Slika 11. Analiza distribucije TAKE / REVIEW / DEFER odluka.",
        [
            "TAKE, REVIEW i DEFER grupe imaju različite profile očekivane vrijednosti. REVIEW je namjerno široka srednja zona u kojoj model ne daje agresivnu preporuku, nego signalizira da je šut blizu replacement benchmarka.",
            "Observed PPS u historijskim grupama je deskriptivan. On ne dokazuje da bi agent kauzalno proizveo te poene, jer nema stvarnog kontrafaktuala za šuteve koji nisu uzeti."
        ],
    )

    add_h1(doc, "11. Eksperimentalni kontekst odbrane")
    add_body(doc, "Produkcijski model koristi moderne NBA podatke od 2020-21 do 2025-26. Defanzivni eksperiment koristi dodatni 2014/15 shot-level tracking dataset od 128.069 šuteva. Taj dataset sadrži korisne varijable koje nisu dostupne u modernom canonical skupu: closest defender distance, shot clock, dribbles i touch time.")
    add_body(doc, "Eksperiment je organizovan kao poređenje više feature grupa. Bazni model koristi shot quality i historiju šutera; zatim se dodaje defensive pressure kroz closest defender distance; zatim offensive creation context kroz shot clock, dribbles i touch time; konačno se ispituje full tracking context. V1.1 model koristi sigurniju konfiguraciju sa tracking kontekstom i shooter rate priorima, bez pokušaja da moderni attempt-count featurei ili defender profile direktno utiču na vjerovatnoću.")
    drows = []
    for label in ["Base shot quality", "Defensive pressure", "Offensive creation context", "Full tracking context"]:
        subset = DEFENSE_CSV[(DEFENSE_CSV["label"] == label) & (DEFENSE_CSV["estimator"] == "hgb")]
        if not subset.empty:
            r = subset.iloc[0]
            drows.append([label, "HGB", fmt(r["test_log_loss"], 5), fmt(r["test_brier"], 5), fmt(r["test_roc_auc"], 5)])
    drows.append(["V1.1 selected (EXPERIMENTAL)", DEFENSE["model"], fmt(DEFENSE["test_log_loss"], 5), fmt(DEFENSE["test_brier"], 5), fmt(DEFENSE["test_roc_auc"], 5)])
    add_table(
        doc,
        "Tabela F. Eksperimentalni defanzivni modeli",
        ["Model / varijanta", "Estimator", "Test Log Loss", "Test Brier", "Test ROC AUC"],
        drows,
    )
    add_figure(
        doc,
        "10_defensive_context_experiment.png",
        "Slika 12. Eksperimentalni uticaj tracking konteksta na Log Loss.",
        [
            "Closest defender distance donosi mjerljiv inkrementalni signal u odnosu na bazni shot quality model. To je košarkaški intuitivno: udaljenost najbližeg defanzivca mijenja kvalitet pokušaja, posebno kod šuteva pod pritiskom.",
            "Ipak, eksperiment ostaje odvojen od produkcijske odluke. Razlog je domain shift između 2014/15 tracking podataka i modernog NBA okruženja, kao i nedostatak savremenog per-shot defender tracking izvora za glavni dataset."
        ],
    )
    add_body(doc, "V1.1 je uveden da bi eksperimentalni prikaz bio sigurniji: koristi robustniju kalibraciju, isključuje pokušaj-count featuree koji nisu semantički bezbjedni za modernu adaptaciju, i uvodi support-aware ponašanje. Kada je scenario izvan podrške trening podataka, aplikacija ne prikazuje apsolutnu eksperimentalnu vjerovatnoću kao pouzdan rezultat.")
    add_body(doc, "Najvažnija akademska lekcija iz ove sekcije je da dodatni kontekst može poboljšati predikciju, ali samo ako je jasno opisano iz kojeg domena dolazi i gdje prestaje njegova pouzdanost. Zbog toga se Advanced Context u aplikaciji prikazuje odvojeno od produkcijskog rezultata, sa jasnim upozorenjem da se radi o eksperimentu.")
    add_body(doc, "Takvo razdvajanje je važno i metodološki i praktično. Da je tracking model direktno uticao na TAKE/REVIEW/DEFER bez modernog ekvivalentnog dataseta, sistem bi miješao pouzdan produkcijski tok sa istraživačkom pretpostavkom. U ovom radu je zato odabrana konzervativnija prezentacija: eksperiment se prikazuje, ali ne mijenja odluku.")

    add_h1(doc, "12. Integracija modela u aplikaciju")
    add_body(doc, "Aplikacija je organizovana kao jednostavan tok od podataka do interaktivne odluke. NBA dataset se obrađuje Python skriptama, model se trenira i serijalizira kroz sklearn/joblib artifacte, Python inference service učitava artefakte, ASP.NET sloj služi kao API/web okruženje, a Decision Lab prikazuje rezultat korisniku.")
    arch = doc.add_paragraph()
    arch.alignment = WD_ALIGN_PARAGRAPH.CENTER
    arch.add_run("NBA datasets → Python preprocessing / sklearn → serialized model artifacts → Python inference service → ASP.NET API proxy → Decision Lab web interface").bold = True
    add_body(doc, "Decision Lab omogućava izbor igrača, lokacije na terenu, vrijednosti šuta, akcije, perioda, vremena, rezultata i home/away konteksta. Produkcijski odgovor sadrži p_make, shot xPPS, replacement xPPS, decision margin i preporuku TAKE/REVIEW/DEFER.")
    add_body(doc, "Glavni endpoint je `POST /api/decision/analyze`. Dodatni endpointi služe za pretragu igrača, profile šutera i historijske decision rezultate. Advanced Context je odvojen eksperimentalni dio koji poziva defanzivni tracking model, ali ne mijenja produkcijsku preporuku.")
    add_body(doc, "Ovakva arhitektura je korisna za akademsku odbranu jer pokazuje da model nije ostao samo u notebooku. Rezultat je povezan sa aplikacijom koja omogućava interaktivno testiranje različitih košarkaških scenarija.")

    add_h1(doc, "13. Ograničenja i mogućnosti poboljšanja")
    add_body(doc, "Najvažnije ograničenje je da su podaci historijski i opservacijski. Model vidi šta se dogodilo u prošlosti, ali nema pravi kontrafaktual: ne zna šta bi se desilo da igrač nije šutirao ili da je lopta završila kod drugog igrača. Zbog toga Decision Policy ne smije biti predstavljen kao dokaz kauzalne optimizacije.")
    add_body(doc, "Drugo ograničenje je playoff drift. Playoff utakmice imaju drugačiji intenzitet, rotacije i selekciju šuteva. Model ostaje koristan za analizu, ali apsolutna kalibracija može odstupati više nego u regularnoj sezoni.")
    add_body(doc, "Treće ograničenje je defanzivni kontekst. Savremeni primarni dataset nema potpuno pouzdan per-shot tracking sa exact defender identity. Zbog toga je defanzivni model treniran na 2014/15 tracking podacima i ostaje eksperimentalan.")
    add_bullets(
        doc,
        [
            "Moguće poboljšanje je nabavka modernog shot-level tracking dataseta za iste sezone kao produkcijski model.",
            "Possession-continuation model mogao bi bolje procijeniti vrijednost odbijenog šuta, ali zahtijeva kvalitetno definisane posjede i kontrafaktualne ciljeve.",
            "Dodatna domain adaptation i playoff kalibracija mogli bi smanjiti drift između regularne sezone i playoffa.",
            "Aplikacija bi se mogla proširiti bazom, autentifikacijom i stabilnim deployment procesom, ali to nije centralni ML cilj ovog rada.",
        ],
    )

    add_h1(doc, "14. Zaključak")
    add_body(doc, "NBA Clutch Agent pokazuje kako se akademski ML projekat može razviti od sirovih sportskih event podataka do funkcionalne decision-support aplikacije. Ključna vrijednost rada nije samo u jednom modelu, nego u metodološkom toku: canonical dataset, temporalna validacija, leakage-safe feature engineering, kalibrisana vjerovatnoća i transparentan decision layer.")
    add_body(doc, "Finalni model na held-out 2025-26 regularnoj sezoni postiže stabilne vjerovatnosne metrike i daje procjene koje se mogu koristiti za očekivane poene po šutu. Decision Policy zatim prevodi tu vjerovatnoću u TAKE, REVIEW i DEFER, uz jasno ograničenje da su historijski observed PPS rezultati deskriptivni, a ne kauzalni dokaz.")
    add_body(doc, "Defanzivni tracking eksperiment pokazuje da dodatni kontekst može biti prediktivno koristan, ali je korektno ostavljen izvan produkcijske odluke zbog domain shift problema i nedostatka modernog ekvivalentnog izvora. Time rad zadržava jasnu granicu između produkcijskog modela i istraživačke ekstenzije.")

    add_h1(doc, "15. Literatura i izvori")
    refs = [
        "NBA Stats, službeni statistički portal: https://www.nba.com/stats",
        "stats.nba.com API i NBA Shot Detail / PlayByPlay izvori korišteni u lokalnim arhivama projekta.",
        "NBA Data Archive korišten za historijske NBA shot i play-by-play podatke u lokalnom workspaceu.",
        "scikit-learn dokumentacija za Logistic Regression, HistGradientBoostingClassifier, kalibraciju i evaluacijske metrike: https://scikit-learn.org/stable/",
        "NBA.com/stats shot logs i tracking/defender kontekst za 2014/15 istraživački defanzivni eksperiment.",
        "Lokalni projektni izvještaji: `reports/data-audit/`, `reports/modeling/`, `reports/decision/` i konsolidovani `submission/metrics_summary.json`.",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="List Number")
        p.add_run(ref)

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
